import sys
import os
import json
import re
import docx.text.paragraph
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def replace_in_para(para, old, new):
    if old not in para.text:
        return
    for run in para.runs:
        if old in run.text:
            run.text = run.text.replace(old, str(new))
            return
    full = para.text
    if old in full:
        new_text = full.replace(old, str(new))
        if para.runs:
            para.runs[0].text = new_text
            for run in para.runs[1:]:
                run.text = ''
        else:
            para.add_run(new_text)

def replace_all(para, replacements):
    for old, new in replacements.items():
        replace_in_para(para, old, str(new))

def safe_text(val, default=''):
    if val is None:
        return default
    if isinstance(val, list):
        return '\n'.join(str(v) for v in val if v)
    return str(val).strip()

def main():
    if len(sys.argv) < 4:
        print('Usage: python labManualGenerator.py <template> <input> <output>')
        sys.exit(1)
    template_path, input_path, output_path = sys.argv[1], sys.argv[2], sys.argv[3]

    with open(input_path, 'r', encoding='utf-8') as f:
        payload = json.load(f)

    subjectCode    = payload.get('subjectCode',   '')
    subjectName    = payload.get('subjectName',   'Academic Subject')
    staffName      = payload.get('staffName',     'Faculty Member')
    departmentName = payload.get('departmentName','Department')
    year           = payload.get('year',          '')
    semester       = payload.get('semester',      '')
    regulation     = payload.get('regulation',    '')
    content_str    = payload.get('content',       '{}')

    try:
        parsed = json.loads(content_str)
    except Exception as e:
        print(f'Content parse error: {e}')
        parsed = {'experiments': []}

    raw_exps = parsed.get('experiments', [])

    def normalise(exp, idx):
        n = exp.get('experimentNo') or exp.get('no') or (idx + 1)
        exp['no'] = exp['experimentNo'] = n
        bl = exp.get('bloomsTaxonomy') or exp.get('bloomsLevel') or 'K3 - Apply'
        exp['bloomsLevel'] = exp['bloomsTaxonomy'] = bl
        proc = exp.get('procedure', [])
        if isinstance(proc, str):
            exp['procedure'] = [s.strip() for s in proc.split('\n') if s.strip()]
        if not exp.get('program1') and exp.get('program'):
            exp['program1'] = exp['program']
        return exp

    experiments = [normalise(e, i) for i, e in enumerate(raw_exps)]
    while len(experiments) < 10:
        n = len(experiments) + 1
        experiments.append(normalise({'no': n, 'experimentNo': n,
            'title': f'Experiment {n}', 'aim': 'To study the concept.',
            'theory': 'Theory.', 'procedure': ['Step 1', 'Step 2'],
            'result': 'Completed.', 'bloomsLevel': 'K3 - Apply', 'coMapping': 'CO5'}, n-1))
    experiments = experiments[:10]

    try:
        sem_num = int(re.sub(r'\D', '', str(semester)))
        even_odd = 'Even' if sem_num % 2 == 0 else 'Odd'
    except Exception:
        sem_num = 0; even_odd = 'Even'

    reg_year = re.sub(r'\D', '', str(regulation)) or str(regulation)

    if not os.path.exists(template_path):
        print(f'Template not found: {template_path}'); sys.exit(1)

    doc = Document(template_path)
    print(f'Template loaded: {len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables')

    cover_map = {
        '{department}': departmentName.upper(),
        '{Course_Code}': subjectCode,
        '{Course_name}': subjectName,
        '{Course_Name}': subjectName,
        '{academic_Year}': year,
        '{Even/Odd}': even_odd,
        '{year}': reg_year,
        '{Regulation}': regulation,
    }

    for para in doc.paragraphs:
        replace_all(para, cover_map)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    replace_all(para, cover_map)

    print('Cover-page placeholders replaced')

    # TOC
    if doc.tables:
        toc = doc.tables[0]
        for idx, exp in enumerate(experiments):
            ri = idx + 1
            if ri < len(toc.rows):
                row = toc.rows[ri]
                nc = len(row.cells)
                col = 2 if nc >= 3 else 1 if nc >= 2 else 0
                cell = row.cells[col]
                for p in cell.paragraphs:
                    for r in p.runs: r.text = ''
                if cell.paragraphs:
                    cell.paragraphs[0].add_run(safe_text(exp.get('title', '')))

    print('TOC populated')

    # Exp header tables
    TITLE_VARS = ['{EXP TITLE}', '{Exp title}', '{EXP title}', '{exp title}']
    exp_ti = 0
    for t_idx, table in enumerate(doc.tables):
        if t_idx == 0: continue
        cell_texts = ' '.join(c.text for row in table.rows for c in row.cells)
        has_ph = any(v.lower() in cell_texts.lower() for v in TITLE_VARS)
        if has_ph and exp_ti < 10:
            exp = experiments[exp_ti]
            title_str = safe_text(exp.get('title', f'Experiment {exp_ti+1}'))
            ex_no_str = str(exp.get('experimentNo', exp_ti+1))
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for v in TITLE_VARS:
                            replace_in_para(p, v, title_str)
                        if 'EX NO:' in p.text.upper() and len(p.text.strip()) < 20:
                            replace_in_para(p, p.text.strip(), f'EX NO: {ex_no_str}')
            exp_ti += 1

    print(f'{exp_ti} experiment header tables filled')

    # Insert content after each experiment table in body XML
    # The template has Tables 1-9 as experiment header tables (9 experiments).
    # We use doc.tables[1:] which gives exactly these tables in document order.
    # Experiment 10 (if no table) gets appended after experiment 9's last content.
    exp_tables_els = [t._element for t in doc.tables[1:]]  # Skip TOC (Table 0)
    num_tbl_exps = len(exp_tables_els)  # typically 9
    print(f'Found {num_tbl_exps} experiment table elements (template has {num_tbl_exps} exp sections)')

    last_exp_last_el = None   # track last inserted element for experiment 10 appending

    for exp_idx in range(10):
        if exp_idx >= len(experiments): break
        exp = experiments[exp_idx]

        # Build sections list
        sections = []
        aim = safe_text(exp.get('aim', ''))
        if aim: sections.append(('AIM:', aim))

        obj = exp.get('objectives', '')
        if isinstance(obj, list):
            obj = '\n'.join(f'{i+1}. {o}' for i,o in enumerate(obj) if o)
        obj = safe_text(obj)
        if obj: sections.append(('OBJECTIVES:', obj))

        theory = safe_text(exp.get('theory', ''))
        if theory: sections.append(('THEORY:', theory))

        sw = safe_text(exp.get('softwareRequired', exp.get('requirements', '')))
        hw = safe_text(exp.get('hardwareRequired', ''))
        if sw or hw:
            rt = (f'Hardware: {hw}\n' if hw else '') + (f'Software: {sw}' if sw else '')
            sections.append(('REQUIREMENTS:', rt.strip()))

        algo = safe_text(exp.get('algorithm', ''))
        if algo: sections.append(('ALGORITHM:', algo))

        proc = exp.get('procedure', [])
        if isinstance(proc, str): proc = [s.strip() for s in proc.split('\n') if s.strip()]
        if proc:
            pt = '\n'.join(f'{i+1}. {s}' for i,s in enumerate(proc))
            sections.append(('PROCEDURE:', pt))

        prog1 = safe_text(exp.get('program1', exp.get('program', '')))
        if prog1: sections.append(('PROGRAM:', prog1))

        s_in  = safe_text(exp.get('sampleInput',  ''))
        s_out = safe_text(exp.get('sampleOutput', ''))
        if s_in or s_out:
            ot = (f'Sample Input:\n{s_in}\n\n' if s_in else '') + (f'Sample Output:\n{s_out}' if s_out else '')
            sections.append(('OBSERVATIONS & OUTPUT:', ot.strip()))

        result = safe_text(exp.get('result', ''))
        if result: sections.append(('RESULT:', result))

        viva = exp.get('vivaQuestions', [])
        if isinstance(viva, list) and viva:
            vlines = []
            for vi, vq in enumerate(viva[:10]):
                if isinstance(vq, dict):
                    vlines.append(f'Q{vi+1}: {vq.get("question","")}')
                    if vq.get('answer'): vlines.append(f'A{vi+1}: {vq["answer"]}')
                else:
                    vlines.append(f'Q{vi+1}: {vq}')
            if vlines: sections.append(('VIVA QUESTIONS:', '\n'.join(vlines)))
        elif isinstance(viva, str) and viva:
            sections.append(('VIVA QUESTIONS:', viva))

        co = safe_text(exp.get('coMapping', ''))
        bl = safe_text(exp.get('bloomsTaxonomy', exp.get('bloomsLevel', '')))
        if co or bl:
            sections.append(('MAPPING:', f'CO Mapping: {co}  |  Bloom Level: {bl}'))

        refs = safe_text(exp.get('references', ''))
        if refs: sections.append(('REFERENCES:', refs))

        if not sections:
            print(f'Exp {exp_idx+1}: no content to insert'); continue

        # Determine insertion anchor
        if exp_idx < len(exp_tables_els):
            # Has a template table: insert AFTER it
            anchor_el = exp_tables_els[exp_idx]
        else:
            # No template table (experiment 10 when template has 9)
            # Add a page break + heading, then content after last experiment's content
            if last_exp_last_el is None:
                print(f'Exp {exp_idx+1}: no anchor found, skipping'); continue

            # Insert page break paragraph
            pb_el = OxmlElement('w:p')
            last_exp_last_el.addnext(pb_el)
            last_exp_last_el = pb_el
            pb_para = docx.text.paragraph.Paragraph(pb_el, doc.part)
            pb_run = pb_para.add_run()
            from docx.oxml import OxmlElement as OxEl
            br = OxEl('w:br')
            br.set(qn('w:type'), 'page')
            pb_run._r.append(br)

            # Insert experiment header row
            ex_no_str = str(exp.get('experimentNo', exp_idx + 1))
            title_str  = safe_text(exp.get('title', f'Experiment {exp_idx+1}'))
            hdr_el = OxmlElement('w:p')
            last_exp_last_el.addnext(hdr_el)
            last_exp_last_el = hdr_el
            hdr_para = docx.text.paragraph.Paragraph(hdr_el, doc.part)
            hdr_para.paragraph_format.space_before = Pt(6)
            hdr_para.paragraph_format.space_after  = Pt(6)
            hr = hdr_para.add_run(f'EX NO: {ex_no_str}   {title_str}')
            hr.bold = True; hr.font.size = Pt(12); hr.font.name = 'Times New Roman'

            anchor_el = last_exp_last_el

        # Insert sections IN ORDER by building forward from anchor
        # Each addnext goes IMMEDIATELY after last_el
        last_el = anchor_el
        for label, content_text in sections:
            lbl_el = OxmlElement('w:p')
            last_el.addnext(lbl_el); last_el = lbl_el
            lbl_para = docx.text.paragraph.Paragraph(lbl_el, doc.part)
            lbl_para.paragraph_format.space_before = Pt(8)
            lbl_para.paragraph_format.space_after  = Pt(2)
            lr = lbl_para.add_run(label)
            lr.bold = True; lr.font.size = Pt(11); lr.font.name = 'Times New Roman'

            for line in content_text.split('\n'):
                c_el = OxmlElement('w:p')
                last_el.addnext(c_el); last_el = c_el
                c_para = docx.text.paragraph.Paragraph(c_el, doc.part)
                c_para.paragraph_format.space_before = Pt(2)
                c_para.paragraph_format.space_after  = Pt(2)
                c_para.paragraph_format.left_indent  = Pt(18)
                cr = c_para.add_run(line.strip())
                cr.font.size = Pt(11); cr.font.name = 'Times New Roman'

        last_exp_last_el = last_el  # remember for next experiment if needed
        print(f'Exp {exp_idx+1} content inserted ({len(sections)} sections)')

    doc.save(output_path)
    print(f'Saved: {output_path}')
    try: os.remove(input_path)
    except Exception: pass
    print('SUCCESS')

if __name__ == '__main__':
    main()
